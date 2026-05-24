"""
Build CWRU_VAE_Presentation.pptx  (14 slides, dark navy theme)
      CWRU_VAE_Poster.pptx         (1 slide A0 portrait, 3-column)
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Cm, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import os, copy

PLOT = "results/plots"
OUT  = "results"

# ── colour palette ────────────────────────────────────────────────────────────
NAVY    = RGBColor(0x0D, 0x1B, 0x3E)   # slide background
MID     = RGBColor(0x11, 0x2D, 0x60)   # section header bg
BLUE    = RGBColor(0x21, 0x96, 0xF3)   # accent
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
LGRAY   = RGBColor(0xB0, 0xBE, 0xC5)   # muted text
GOLD    = RGBColor(0xFF, 0xD6, 0x00)   # highlight number
GREEN   = RGBColor(0x00, 0xE6, 0x76)   # success metric


# ══════════════════════════════════════════════════════════════════════════════
#  HELPERS
# ══════════════════════════════════════════════════════════════════════════════

def rgb(r, g, b):   return RGBColor(r, g, b)

def add_rect(slide, x, y, w, h, fill=None, line=None):
    from pptx.util import Inches
    shape = slide.shapes.add_shape(
        1,                          # MSO_SHAPE_TYPE.RECTANGLE = 1
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    sf = shape.fill
    if fill:
        sf.solid(); sf.fore_color.rgb = fill
    else:
        sf.background()
    shape.line.fill.background() if not line else None
    return shape


def add_txt(slide, text, x, y, w, h, size=14, bold=False, color=WHITE,
            align=PP_ALIGN.LEFT, italic=False, wrap=True):
    txb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf  = txb.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return txb


def add_img(slide, path, x, y, w, h=None):
    if h:
        slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w))


def dark_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY


def section_bar(slide, title, y=0.0):
    """Dark blue title bar across full width."""
    add_rect(slide, 0, y, 10, 0.52, fill=MID)
    add_txt(slide, title, 0.18, y+0.04, 9.6, 0.44,
            size=20, bold=True, color=BLUE, align=PP_ALIGN.LEFT)


def metric_card(slide, value, label, x, y, w=2.2, h=0.9,
                val_color=GOLD, bg=MID):
    add_rect(slide, x, y, w, h, fill=bg)
    add_txt(slide, value, x+0.08, y+0.04, w-0.16, 0.45,
            size=22, bold=True, color=val_color, align=PP_ALIGN.CENTER)
    add_txt(slide, label, x+0.08, y+0.48, w-0.16, 0.38,
            size=9, color=LGRAY, align=PP_ALIGN.CENTER)


def bullet_block(slide, items, x, y, w, size=11, color=WHITE, gap=0.28):
    for i, item in enumerate(items):
        add_txt(slide, f"▸  {item}", x, y + i*gap, w, gap+0.04,
                size=size, color=color)


# ══════════════════════════════════════════════════════════════════════════════
#  PRESENTATION  (14 slides, 10" × 5.625")
# ══════════════════════════════════════════════════════════════════════════════

def make_presentation():
    prs = Presentation()
    prs.slide_width  = Inches(10)
    prs.slide_height = Inches(5.625)
    blank = prs.slide_layouts[6]   # completely blank

    # ── Slide 1 : Title ──────────────────────────────────────────────────────
    s = prs.slides.add_slide(blank); dark_bg(s)
    add_rect(s, 0, 0, 10, 1.1, fill=MID)
    add_txt(s, "CWRU Bearing Dataset  |  VAE  |  Anomaly Detection",
            0.3, 0.18, 9.4, 0.5, size=10, color=LGRAY, align=PP_ALIGN.CENTER)
    add_txt(s,
        "Unsupervised Bearing Fault Detection\nUsing Variational Autoencoder",
        0.4, 1.25, 9.2, 1.6, size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_txt(s,
        "KL-Regularised Latent Normal Manifold for Industrial Anomaly Detection",
        0.4, 2.9, 9.2, 0.55, size=14, color=BLUE, italic=True, align=PP_ALIGN.CENTER)
    metric_card(s, "AUC = 1.0000", "Perfect Detection",   1.0, 4.1, val_color=GREEN)
    metric_card(s, "5-Fold CV",    "Cross Validation",    3.4, 4.1)
    metric_card(s, "5.2×",         "Fault/Normal Error",  5.8, 4.1, val_color=GOLD)
    metric_card(s, "Normal Only",  "Training Strategy",   8.0, 4.1, val_color=BLUE,
                w=1.8, bg=rgb(0x0D,0x47,0xA1))

    # ── Slide 2 : Motivation ─────────────────────────────────────────────────
    s = prs.slides.add_slide(blank); dark_bg(s)
    section_bar(s, "Background & Motivation")
    add_txt(s, "Why Bearing Fault Detection Matters", 0.3, 0.65, 5.5, 0.4,
            size=14, bold=True, color=BLUE)
    bullet_block(s, [
        "Bearings are critical in rotating machinery (motors, turbines, pumps)",
        "Undetected faults → unexpected downtime, safety risks, economic loss",
        "Early detection can reduce maintenance costs by up to 30%",
        "Industrial IoT demands real-time, automated monitoring solutions",
    ], 0.3, 1.1, 5.5, size=11)

    add_txt(s, "The Core Challenge", 0.3, 2.55, 5.5, 0.4,
            size=14, bold=True, color=BLUE)
    bullet_block(s, [
        "Fault events are RARE → severe class imbalance",
        "Labelling fault data is costly & requires expert knowledge",
        "Traditional supervised methods fail under imbalanced conditions",
    ], 0.3, 3.0, 5.5, size=11)

    add_txt(s, "Our Solution", 0.3, 4.05, 5.5, 0.35,
            size=14, bold=True, color=GREEN)
    add_txt(s, "Train VAE exclusively on normal data → learn the 'normal manifold'\n"
               "→ Any signal outside this manifold is flagged as anomalous",
            0.3, 4.4, 5.5, 0.8, size=11, color=WHITE)

    add_rect(s, 5.9, 0.65, 3.8, 4.6, fill=MID)
    add_txt(s, "Key Insight", 6.0, 0.75, 3.6, 0.4,
            size=13, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    for label, val, yy, vc in [
        ("Total Segments",   "11,832",          1.3,  WHITE),
        ("After Balancing",  "4,162",           1.85, WHITE),
        ("Normal (Train)",   "3,310  (79.5%)",  2.4,  GREEN),
        ("Fault (Eval only)","852   (20.5%)",   2.95, LGRAY),
        ("Inner Race",       "282 segments",    3.45, rgb(0xF4,0x43,0x36)),
        ("Ball Fault",       "271 segments",    3.9,  rgb(0xFF,0x98,0x00)),
        ("Outer Race",       "299 segments",    4.35, rgb(0x9C,0x27,0xB0)),
    ]:
        add_txt(s, label, 6.0, yy, 2.2, 0.38, size=10, color=LGRAY)
        add_txt(s, val,   8.0, yy, 1.7, 0.38, size=10, bold=True, color=vc, align=PP_ALIGN.RIGHT)

    # ── Slide 3 : Dataset ────────────────────────────────────────────────────
    s = prs.slides.add_slide(blank); dark_bg(s)
    section_bar(s, "CWRU Bearing Dataset")
    add_img(s, f"{PLOT}/01_time_domain.png", 5.2, 0.6, 4.7, 4.7)

    rows = [
        ("Bearing Model",    "SKF 6205 Deep Groove Ball"),
        ("Sampling Rate",    "12,000 Hz"),
        ("Fault Method",     "Electro-Discharge Machining (EDM)"),
        ("Fault Sizes",      "0.007\" / 0.014\" / 0.021\" dia."),
        ("Motor Loads",      "0, 1, 2, 3 HP  (1720–1797 RPM)"),
        ("Total Files",      "40 .mat files (4 Normal + 36 Fault)"),
    ]
    add_txt(s, "Experimental Setup", 0.3, 0.65, 4.7, 0.35,
            size=13, bold=True, color=BLUE)
    for i, (k, v) in enumerate(rows):
        yy = 1.1 + i * 0.52
        add_rect(s, 0.3, yy, 4.6, 0.44,
                 fill=rgb(0x11,0x2D,0x60) if i%2==0 else rgb(0x0D,0x1B,0x3E))
        add_txt(s, k, 0.4, yy+0.06, 1.9, 0.32, size=10, color=LGRAY)
        add_txt(s, v, 2.3, yy+0.06, 2.5, 0.32, size=10, bold=True, color=WHITE)

    add_txt(s, "Fault Classes", 0.3, 4.2, 4.7, 0.3, size=12, bold=True, color=BLUE)
    for i, (name, col) in enumerate([
        ("Normal",          GREEN),
        ("Inner Race Fault",rgb(0xF4,0x43,0x36)),
        ("Ball Fault",      rgb(0xFF,0x98,0x00)),
        ("Outer Race Fault",rgb(0x9C,0x27,0xB0)),
    ]):
        add_rect(s, 0.3+i*1.18, 4.55, 1.1, 0.65, fill=MID)
        add_txt(s, name, 0.32+i*1.18, 4.6, 1.06, 0.55,
                size=8, bold=True, color=col, align=PP_ALIGN.CENTER)

    # ── Slide 4 : Preprocessing ──────────────────────────────────────────────
    s = prs.slides.add_slide(blank); dark_bg(s)
    section_bar(s, "Data Preprocessing & Factory-Like Balancing")
    add_img(s, f"{PLOT}/04_class_distribution.png", 4.8, 0.6, 5.0, 4.7)

    steps = [
        ("1  Sliding Window",
         "Size: 1,024 samples (≈85 ms @ 12 kHz)\nStride: 512 samples (50% overlap)"),
        ("2  Factory-Like Balancing",
         "Keep ALL normal segments (3,310)\nRandom 10% of fault segments → 852\nSimulates real plant: faults are rare"),
        ("3  Z-Score Normalisation",
         "Per-sample: subtract mean, divide by std\nEnsures equal amplitude scale across windows"),
        ("4  Train / Eval Split",
         "VAE trains on normal only → 3,310 segments\nFault data used ONLY for evaluation (AUC)"),
    ]
    y = 0.65
    for title, body in steps:
        add_rect(s, 0.25, y, 4.4, 1.1, fill=MID)
        add_txt(s, title, 0.35, y+0.06, 4.2, 0.36, size=11, bold=True, color=BLUE)
        add_txt(s, body,  0.35, y+0.42, 4.2, 0.62, size=9,  color=WHITE)
        y += 1.18

    # ── Slide 5 : Signal Analysis ────────────────────────────────────────────
    s = prs.slides.add_slide(blank); dark_bg(s)
    section_bar(s, "Signal Analysis — FFT Spectrum & STFT Spectrogram")
    add_img(s, f"{PLOT}/02_fft_spectrum.png", 0.1, 0.6, 4.9, 4.7)
    add_img(s, f"{PLOT}/03_spectrogram.png",  5.1, 0.6, 4.8, 4.7)
    add_txt(s, "Frequency Domain (FFT)",  0.2, 5.2, 4.7, 0.35,
            size=11, color=LGRAY, align=PP_ALIGN.CENTER)
    add_txt(s, "Time-Frequency (STFT)",   5.2, 5.2, 4.7, 0.35,
            size=11, color=LGRAY, align=PP_ALIGN.CENTER)

    # ── Slide 6 : t-SNE ──────────────────────────────────────────────────────
    s = prs.slides.add_slide(blank); dark_bg(s)
    section_bar(s, "t-SNE Visualisation of Raw Signal Feature Space")
    add_img(s, f"{PLOT}/05_tsne_features.png", 0.5, 0.65, 9.0, 4.6)
    add_txt(s,
        "Fault signals are partially separable from normal in raw space — VAE latent space achieves complete separation.",
        0.5, 5.25, 9.0, 0.35, size=10, color=LGRAY, italic=True, align=PP_ALIGN.CENTER)

    # ── Slide 7 : VAE Architecture ───────────────────────────────────────────
    s = prs.slides.add_slide(blank); dark_bg(s)
    section_bar(s, "VAE Architecture — 1-D Convolutional Encoder–Decoder")

    # Encoder box
    add_rect(s, 0.2, 0.65, 3.0, 4.5, fill=MID)
    add_txt(s, "ENCODER", 0.3, 0.72, 2.8, 0.38, size=13, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    enc_layers = [
        "Input  (B, 1, 1024)",
        "Conv1d 1→16  k7 s2",
        "Conv1d 16→32 k5 s2",
        "Conv1d 32→64 k3 s2",
        "Conv1d 64→128 k3 s2",
        "Flatten → 8,192",
        "FC → 256  (ReLU)",
        "FC → μ (32)   FC → logσ² (32)",
    ]
    for i, l in enumerate(enc_layers):
        add_txt(s, l, 0.3, 1.15+i*0.38, 2.8, 0.34, size=9,
                color=GOLD if "μ" in l else WHITE, align=PP_ALIGN.CENTER)

    # Middle
    add_rect(s, 3.4, 0.65, 3.2, 4.5, fill=rgb(0x0A,0x23,0x4F))
    add_txt(s, "LATENT SPACE  z ∈ ℝ³²", 3.5, 0.75, 3.0, 0.4,
            size=12, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    add_txt(s, "Reparameterisation Trick", 3.5, 1.3, 3.0, 0.3,
            size=10, bold=True, color=LGRAY, align=PP_ALIGN.CENTER)
    add_txt(s, "z = μ + ε · exp(0.5·logσ²)\nε ~ N(0, I)",
            3.5, 1.65, 3.0, 0.6, size=11, color=WHITE, align=PP_ALIGN.CENTER, italic=True)
    add_rect(s, 3.55, 2.5, 2.9, 1.5, fill=rgb(0x0D,0x1B,0x3E))
    add_txt(s, "Loss Function:", 3.65, 2.58, 2.7, 0.3,
            size=10, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    add_txt(s, "L = MSE(x, x̂)×1024\n   + β · KL(q ‖ N(0,I))",
            3.65, 2.92, 2.7, 0.6, size=10, color=WHITE, italic=True, align=PP_ALIGN.CENTER)
    add_txt(s, "β = 1.0  (Standard VAE)", 3.65, 3.6, 2.7, 0.3,
            size=9, color=LGRAY, align=PP_ALIGN.CENTER)
    add_txt(s, "Anomaly Score", 3.55, 4.2, 2.9, 0.3,
            size=10, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    add_txt(s, "= per-sample MSE(x, x̂)\nHigh score → Anomaly",
            3.55, 4.5, 2.9, 0.5, size=9, color=WHITE, align=PP_ALIGN.CENTER)

    # Decoder box
    add_rect(s, 6.8, 0.65, 3.0, 4.5, fill=MID)
    add_txt(s, "DECODER", 6.9, 0.72, 2.8, 0.38, size=13, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    dec_layers = [
        "z  (B, 32)",
        "FC → 256  (ReLU)",
        "FC → 8,192  (ReLU)",
        "Reshape  (128, 64)",
        "ConvT 128→64 k3 s2",
        "ConvT 64→32  k3 s2",
        "ConvT 32→16  k5 s2",
        "ConvT 16→1   k7 s2  + Tanh",
    ]
    for i, l in enumerate(dec_layers):
        add_txt(s, l, 6.9, 1.15+i*0.38, 2.8, 0.34, size=9,
                color=GREEN if i==7 else WHITE, align=PP_ALIGN.CENTER)

    # ── Slide 8 : Training Strategy ───────────────────────────────────────────
    s = prs.slides.add_slide(blank); dark_bg(s)
    section_bar(s, "Training Strategy — Normal-Only Unsupervised Learning")
    cards = [
        ("Normal-Only Training", "VAE sees ONLY normal bearing signals\nduring training — no fault labels needed"),
        ("5-Fold Cross Validation", "KFold(n=5, shuffle=True)\nEach fold: 2,648 train / 662 val"),
        ("Optimizer",  "Adam  lr=1e-3  weight_decay=1e-5\nReduceLROnPlateau(patience=10, f=0.5)"),
        ("Early Stopping", "Patience = 20 epochs\nRestore best val-loss checkpoint"),
        ("Batch Size / Epochs", "Batch = 64,  Max epochs = 100\nFold 5 stopped at epoch 60"),
        ("Hardware", "Apple MPS (Metal Performance Shaders)\nMacBook Apple Silicon GPU acceleration"),
    ]
    for i, (t, b) in enumerate(cards):
        col = i % 3; row = i // 3
        x = 0.2 + col * 3.27; y = 0.7 + row * 2.3
        add_rect(s, x, y, 3.1, 2.1, fill=MID)
        add_txt(s, t, x+0.12, y+0.1, 2.86, 0.4, size=12, bold=True, color=BLUE)
        add_txt(s, b, x+0.12, y+0.55, 2.86, 1.3, size=10, color=WHITE)

    # ── Slide 9 : Training Curves ─────────────────────────────────────────────
    s = prs.slides.add_slide(blank); dark_bg(s)
    section_bar(s, "VAE Training Curves — Best Fold (Fold 1)")
    add_img(s, f"{PLOT}/07_training_curves.png", 0.15, 0.62, 9.7, 4.65)
    add_txt(s,
        "Total loss: 1,172 → 176  |  Reconstruction: 1,172 → 173  |  KL divergence stabilises at 3.79",
        0.2, 5.3, 9.6, 0.3, size=10, color=LGRAY, italic=True, align=PP_ALIGN.CENTER)

    # ── Slide 10 : K-Fold ─────────────────────────────────────────────────────
    s = prs.slides.add_slide(blank); dark_bg(s)
    section_bar(s, "5-Fold Cross-Validation Results")
    add_img(s, f"{PLOT}/06_kfold_results.png", 0.15, 0.62, 9.7, 4.3)
    metric_card(s, "1.0000 ± 0.0000", "Mean AUC-ROC (all 5 folds)", 1.5, 5.05, w=3.2, val_color=GREEN)
    metric_card(s, "342.61 ± 3.72",  "Mean Val Loss",               5.3, 5.05, w=3.2)

    # ── Slide 11 : Reconstruction Error ──────────────────────────────────────
    s = prs.slides.add_slide(blank); dark_bg(s)
    section_bar(s, "Anomaly Detection — Reconstruction Error Analysis")
    add_img(s, f"{PLOT}/08_reconstruction_error.png", 0.15, 0.62, 9.7, 4.2)
    for val, lbl, x, vc in [
        ("0.2007", "Normal  μ",  0.4,  BLUE),
        ("0.0950", "Normal  σ",  2.1,  LGRAY),
        ("1.0493", "Fault   μ",  3.8,  rgb(0xF4,0x43,0x36)),
        ("0.0103", "Fault   σ",  5.5,  LGRAY),
        ("5.2×",   "Error Ratio",7.2,  GOLD),
        ("0.4778", "Threshold",  8.9,  GREEN),
    ]:
        add_rect(s, x, 4.95, 1.55, 0.6, fill=MID)
        add_txt(s, val, x+0.05, 4.98, 1.45, 0.3, size=14, bold=True, color=vc, align=PP_ALIGN.CENTER)
        add_txt(s, lbl, x+0.05, 5.27, 1.45, 0.25, size=8,  color=LGRAY, align=PP_ALIGN.CENTER)

    # ── Slide 12 : ROC ────────────────────────────────────────────────────────
    s = prs.slides.add_slide(blank); dark_bg(s)
    section_bar(s, "ROC Curve — VAE Anomaly Detector")
    add_img(s, f"{PLOT}/09_roc_curve.png", 1.0, 0.62, 8.0, 4.7)
    add_txt(s, "AUC = 1.0000  — Perfect separation between normal and fault signals",
            0.4, 5.35, 9.2, 0.3, size=11, bold=True, color=GREEN, align=PP_ALIGN.CENTER)

    # ── Slide 13 : Latent Space ───────────────────────────────────────────────
    s = prs.slides.add_slide(blank); dark_bg(s)
    section_bar(s, "VAE Latent Space — t-SNE Projection")
    add_img(s, f"{PLOT}/10_latent_space.png", 0.15, 0.62, 9.7, 4.45)
    add_txt(s,
        "Normal data forms a diffuse ring (learned N(0,I) prior).  "
        "Fault data clusters centrally — outside the normal reconstruction manifold.",
        0.4, 5.12, 9.2, 0.45, size=10, color=LGRAY, italic=True, align=PP_ALIGN.CENTER)

    # ── Slide 14 : Conclusion ─────────────────────────────────────────────────
    s = prs.slides.add_slide(blank); dark_bg(s)
    section_bar(s, "Conclusion & Future Work")
    add_txt(s, "Key Achievements", 0.3, 0.65, 5.5, 0.35, size=14, bold=True, color=BLUE)
    bullet_block(s, [
        "AUC-ROC = 1.0000 across all 5 cross-validation folds",
        "Fault reconstruction error 5.2× higher than normal",
        "Threshold 0.4778 cleanly separates all fault types",
        "Unsupervised — no fault labels required for training",
        "Factory-realistic distribution: 79.5% normal data",
    ], 0.3, 1.05, 5.5)

    add_txt(s, "Summary Results", 5.8, 0.65, 3.9, 0.35, size=14, bold=True, color=BLUE)
    for label, val, yy, vc in [
        ("K-Fold AUC",         "1.0000 ± 0.000", 1.08, GREEN),
        ("Final AUC",          "1.0000",           1.58, GREEN),
        ("Normal Error μ",     "0.2007",           2.08, BLUE),
        ("Fault Error μ",      "1.0493",           2.58, rgb(0xF4,0x43,0x36)),
        ("Error Ratio",        "5.2×",             3.08, GOLD),
        ("Threshold (P95)",    "0.4778",           3.58, GREEN),
    ]:
        add_rect(s, 5.8, yy, 3.8, 0.44, fill=MID)
        add_txt(s, label, 5.9, yy+0.06, 2.0, 0.32, size=10, color=LGRAY)
        add_txt(s, val,   7.8, yy+0.06, 1.7, 0.32, size=10, bold=True, color=vc, align=PP_ALIGN.RIGHT)

    add_txt(s, "Future Work", 0.3, 4.15, 9.2, 0.35, size=14, bold=True, color=BLUE)
    bullet_block(s, [
        "Deploy on streaming IoT sensor data with real-time threshold adaptation",
        "Extend to multi-sensor fusion (vibration + temperature + current)",
        "Integrate transformer-based encoder for long-range temporal dependencies",
    ], 0.3, 4.5, 9.4, size=10, gap=0.3)

    out_path = f"{OUT}/CWRU_VAE_Presentation.pptx"
    prs.save(out_path)
    print(f"  Saved → {out_path}  ({len(prs.slides)} slides)")


# ══════════════════════════════════════════════════════════════════════════════
#  POSTER  (1 slide, A0 portrait 84 cm × 119 cm)
# ══════════════════════════════════════════════════════════════════════════════

def make_poster():
    # A0 portrait in inches: 33.07 × 46.81
    W, H = 33.07, 46.81

    prs = Presentation()
    prs.slide_width  = Inches(W)
    prs.slide_height = Inches(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    # White background
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE

    # ── Top header ───────────────────────────────────────────────────────────
    add_rect(s, 0, 0, W, 3.5, fill=rgb(0x0D,0x47,0xA1))
    add_txt(s, "Unsupervised Bearing Fault Detection Using Variational Autoencoder",
            0.4, 0.25, W-0.8, 1.6, size=48, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_txt(s, "KL-Regularised Latent Normal Manifold · CWRU Dataset · AUC = 1.0000",
            0.4, 1.9, W-0.8, 0.9, size=26, color=rgb(0xBB,0xDE,0xFB), italic=True, align=PP_ALIGN.CENTER)
    add_txt(s, "VAE · Conv1D · 5-Fold CV · Apple MPS",
            0.4, 2.75, W-0.8, 0.65, size=20, color=rgb(0x90,0xCA,0xF9), align=PP_ALIGN.CENTER)

    # ── 3 column layout ──────────────────────────────────────────────────────
    col_w = 10.2
    col_x = [0.4, 11.4, 22.4]
    y0    = 3.8
    GAP   = 0.25

    def ph(slide, title, x, y, w, col=rgb(0x0D,0x47,0xA1)):
        """Section header for poster."""
        add_rect(slide, x, y, w, 0.65, fill=col)
        add_txt(slide, title, x+0.15, y+0.08, w-0.3, 0.5,
                size=18, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
        return y + 0.65 + 0.15

    def pt(slide, text, x, y, w, size=13, color=rgb(0x1A,0x1A,0x2E)):
        add_txt(slide, text, x, y, w, 6, size=size, color=color, wrap=True)

    DARK = rgb(0x1A, 0x1A, 0x2E)

    # ════ COLUMN 1 ════════════════════════════════════════════════════════════
    cx, cw = col_x[0], col_w
    y = y0

    y = ph(s, "1. Introduction & Motivation", cx, y, cw)
    intro = (
        "Rolling element bearings are critical in industrial machinery. "
        "Undetected faults cause costly unplanned downtime. "
        "Real plants exhibit severely imbalanced data — faults are rare events.\n\n"
        "Traditional supervised fault diagnosis requires labelled fault data, "
        "which is expensive to collect and often unavailable.\n\n"
        "We propose an unsupervised VAE trained exclusively on normal "
        "vibration signals, learning a compact latent manifold N(0,I). "
        "Fault signals cannot be well-reconstructed → high reconstruction error "
        "→ anomaly flag."
    )
    pt(s, intro, cx, y, cw); y += 2.6 + GAP

    y = ph(s, "2. CWRU Bearing Dataset", cx, y, cw)
    tbl_data = [
        ["Parameter",       "Value"],
        ["Bearing Model",   "SKF 6205"],
        ["Sampling Rate",   "12,000 Hz"],
        ["Fault Method",    "EDM (single point)"],
        ["Fault Sizes",     "0.007 / 0.014 / 0.021\""],
        ["Motor Loads",     "0–3 HP (1720–1797 RPM)"],
        ["Total Files",     "40 .mat files"],
        ["Fault Classes",   "IR / Ball / Outer Race"],
    ]
    tbl = s.shapes.add_table(len(tbl_data), 2,
        Inches(cx), Inches(y), Inches(cw), Inches(3.6)).table
    for r, row in enumerate(tbl_data):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            tf = cell.text_frame
            tf.paragraphs[0].font.size = Pt(12)
            tf.paragraphs[0].font.bold = (r==0)
            tf.paragraphs[0].font.color.rgb = WHITE if r==0 else DARK
            if r == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb(0x0D,0x47,0xA1)
            elif r%2 == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb(0xE8,0xF0,0xFE)
    y += 3.7 + GAP

    y = ph(s, "3. Data Preprocessing", cx, y, cw)
    pt(s, "• Sliding window: 1,024 samples (≈85 ms), stride 512 (50% overlap)\n"
          "• Factory-like balancing: all normal (3,310) + 10% fault (852)\n"
          "• Total balanced dataset: 4,162 segments\n"
          "• Per-sample Z-score normalisation",
       cx, y, cw); y += 1.6 + GAP
    add_img(s, f"{PLOT}/04_class_distribution.png", cx, y, cw); y += 5.8 + GAP

    y = ph(s, "4. Time-Domain Signals", cx, y, cw)
    add_img(s, f"{PLOT}/01_time_domain.png", cx, y, cw)

    # ════ COLUMN 2 ════════════════════════════════════════════════════════════
    cx, cw = col_x[1], col_w
    y = y0

    y = ph(s, "5. VAE Architecture", cx, y, cw)
    arch = (
        "ENCODER\n"
        "  Input (B, 1, 1024)\n"
        "  Conv1d 1→16  k=7 s=2  BN  ReLU\n"
        "  Conv1d 16→32 k=5 s=2  BN  ReLU\n"
        "  Conv1d 32→64 k=3 s=2  BN  ReLU\n"
        "  Conv1d 64→128 k=3 s=2 BN  ReLU\n"
        "  Flatten → 8,192\n"
        "  FC 8192→256 ReLU\n"
        "  FC 256 → μ(32)  &  FC 256 → logσ²(32)\n\n"
        "REPARAMETERISATION\n"
        "  z = μ + ε·exp(0.5·logσ²),  ε ~ N(0,I)\n\n"
        "DECODER (mirror of Encoder)\n"
        "  z(32) → FC → ConvTranspose×4 → Output(B,1,1024)\n\n"
        "LOSS FUNCTION\n"
        "  L = MSE(x,x̂)×1024  +  β·KL(q ‖ N(0,I))\n"
        "  KL = -½·E[1 + logσ² - μ² - σ²]\n"
        "  β = 1.0  (Standard VAE)"
    )
    pt(s, arch, cx, y, cw, size=12); y += 6.8 + GAP

    y = ph(s, "6. Training Configuration", cx, y, cw)
    cfg = (
        "• Train ONLY on normal signals (unsupervised)\n"
        "• 5-Fold Cross-Validation (KFold, shuffle=True)\n"
        "• Optimizer: Adam  lr=1e-3  weight_decay=1e-5\n"
        "• LR Scheduler: ReduceLROnPlateau (patience=10)\n"
        "• Early stopping: patience=20  (Fold 5 → epoch 60)\n"
        "• Batch size: 64  |  Max epochs: 100\n"
        "• Device: Apple MPS (Metal Performance Shaders)"
    )
    pt(s, cfg, cx, y, cw, size=12); y += 2.5 + GAP

    y = ph(s, "7. FFT Frequency Spectrum", cx, y, cw)
    add_img(s, f"{PLOT}/02_fft_spectrum.png", cx, y, cw); y += 6.5 + GAP

    y = ph(s, "8. STFT Spectrogram", cx, y, cw)
    add_img(s, f"{PLOT}/03_spectrogram.png", cx, y, cw); y += 5.5 + GAP

    y = ph(s, "9. t-SNE Feature Space", cx, y, cw)
    add_img(s, f"{PLOT}/05_tsne_features.png", cx, y, cw)

    # ════ COLUMN 3 ════════════════════════════════════════════════════════════
    cx, cw = col_x[2], col_w
    y = y0

    y = ph(s, "10. Training Curves (Best Fold)", cx, y, cw)
    add_img(s, f"{PLOT}/07_training_curves.png", cx, y, cw); y += 4.5 + GAP

    y = ph(s, "11. K-Fold Cross-Validation", cx, y, cw)
    add_img(s, f"{PLOT}/06_kfold_results.png", cx, y, cw); y += 5.3 + GAP

    y = ph(s, "12. Reconstruction Error Distribution", cx, y, cw)
    add_img(s, f"{PLOT}/08_reconstruction_error.png", cx, y, cw); y += 5.3 + GAP

    y = ph(s, "13. ROC Curve", cx, y, cw)
    add_img(s, f"{PLOT}/09_roc_curve.png", cx, y, cw); y += 5.5 + GAP

    y = ph(s, "14. Latent Space (t-SNE)", cx, y, cw)
    add_img(s, f"{PLOT}/10_latent_space.png", cx, y, cw); y += 5.5 + GAP

    y = ph(s, "15. Results & Conclusion", cx, y, cw, col=rgb(0x1B,0x5E,0x20))
    results = [
        ("K-Fold AUC-ROC",     "1.0000 ± 0.0000", GREEN),
        ("Final AUC-ROC",      "1.0000",           GREEN),
        ("Normal Error (μ±σ)", "0.2007 ± 0.0950",  rgb(0x21,0x96,0xF3)),
        ("Fault Error  (μ±σ)", "1.0493 ± 0.0103",  rgb(0xF4,0x43,0x36)),
        ("Error Ratio",        "~5.2×",             rgb(0xFF,0xD6,0x00)),
        ("Threshold (P95)",    "0.4778",            GREEN),
    ]
    for label, val, vc in results:
        add_rect(s, cx, y, cw, 0.7, fill=rgb(0xE8,0xF5,0xE9))
        add_txt(s, label, cx+0.1, y+0.1, 6.5, 0.5, size=13, color=DARK)
        add_txt(s, val,   cx+6.6, y+0.1, 3.5, 0.5, size=14, bold=True, color=vc,
                align=PP_ALIGN.RIGHT)
        y += 0.75

    y += 0.3
    pt(s, "Conclusion: The VAE trained exclusively on normal bearing vibration "
          "signals achieves perfect anomaly detection (AUC=1.000) across all "
          "5 cross-validation folds. The KL regularisation enforces a compact "
          "normal manifold, making fault signals easily identifiable by their "
          "5.2× higher reconstruction error. This approach is directly applicable "
          "to real industrial monitoring where fault labels are unavailable.",
       cx, y, cw, size=12)

    # ── Footer ────────────────────────────────────────────────────────────────
    add_rect(s, 0, H-1.0, W, 1.0, fill=rgb(0x0D,0x47,0xA1))
    add_txt(s, "Dataset: Case Western Reserve University Bearing Data Center  "
               "·  Model: 1-D Conv VAE  ·  Framework: PyTorch  ·  Device: Apple MPS",
            0.4, H-0.85, W-0.8, 0.7, size=14, color=rgb(0x90,0xCA,0xF9),
            align=PP_ALIGN.CENTER)

    out_path = f"{OUT}/CWRU_VAE_Poster.pptx"
    prs.save(out_path)
    print(f"  Saved → {out_path}  (A0 portrait poster)")


# ══════════════════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    os.chdir("/Users/guset/deep/cwru_vae")
    print("Building Presentation …")
    make_presentation()
    print("Building Poster …")
    make_poster()
    print("Done.")
