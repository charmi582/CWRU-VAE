from pathlib import Path

import pandas as pd
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "presentation"
ASSET_DIR = OUT_DIR / "assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

load_df = pd.read_csv(ROOT / "results/load_wise/load_wise_metrics.csv")
cal_df = pd.read_csv(ROOT / "results/threshold_calibration/threshold_calibration_metrics.csv")
noise_df = pd.read_csv(ROOT / "results/noise_robustness/noise_robustness_metrics.csv")
ae_df = pd.read_csv(ROOT / "results/ae_load_wise/ae_load_wise_metrics.csv")

cal_order = [
    "train_p95",
    "val_p95",
    "val_p99",
    "target_load_10pct_normal_p95",
    "target_load_10pct_normal_p99",
]
cal_summary = (
    cal_df.groupby("strategy")[["f1", "false_alarm_rate"]]
    .mean()
    .reset_index()
    .set_index("strategy")
    .loc[cal_order]
    .reset_index()
)
noise_order = ["clean", "30", "20", "10", "5"]
noise_summary = noise_df.groupby("snr_db")[["f1", "false_alarm_rate"]].mean().reset_index()
noise_summary["snr_db"] = noise_summary["snr_db"].astype(str)
noise_summary = noise_summary.set_index("snr_db").loc[noise_order].reset_index()

plt.rcParams["font.sans-serif"] = [
    "Microsoft JhengHei",
    "Noto Sans CJK TC",
    "Arial Unicode MS",
    "DejaVu Sans",
]
plt.rcParams["axes.unicode_minus"] = False


def save_bar_chart(path, labels, vals1, vals2=None, title="", ylabel="Percent (%)", leg1="F1", leg2="FAR", ylim=(0, 105)):
    fig, ax = plt.subplots(figsize=(8.2, 4.2), dpi=180)
    x = list(range(len(labels)))
    if vals2 is None:
        ax.bar(x, vals1, color="#0F766E", width=0.58)
    else:
        w = 0.36
        ax.bar([i - w / 2 for i in x], vals1, width=w, label=leg1, color="#0F766E")
        ax.bar([i + w / 2 for i in x], vals2, width=w, label=leg2, color="#B45309")
        ax.legend(frameon=False, loc="upper left")
    ax.set_title(title, fontsize=15, weight="bold", pad=12)
    ax.set_ylabel(ylabel)
    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    ax.set_ylim(*ylim)
    ax.grid(axis="y", color="#E5E7EB", linewidth=0.8)
    for spine in ["top", "right"]:
        ax.spines[spine].set_visible(False)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


save_bar_chart(
    ASSET_DIR / "loadwise_far.png",
    [str(x) for x in load_df["heldout_load"]],
    (load_df["false_alarm_rate"] * 100).tolist(),
    title="Load-wise split: fixed P95 threshold exposes FAR shift",
    ylabel="False alarm rate (%)",
)
save_bar_chart(
    ASSET_DIR / "calibration.png",
    ["Train P95", "Val P95", "Val P99", "Target 10% P95", "Target 10% P99"],
    (cal_summary["f1"] * 100).tolist(),
    (cal_summary["false_alarm_rate"] * 100).tolist(),
    title="Threshold calibration improves deployability",
    leg1="Mean F1",
    leg2="Mean FAR",
)
save_bar_chart(
    ASSET_DIR / "noise.png",
    ["Clean", "30 dB", "20 dB", "10 dB", "5 dB"],
    (noise_summary["f1"] * 100).tolist(),
    (noise_summary["false_alarm_rate"] * 100).tolist(),
    title="Noise robustness: fixed threshold fails at low SNR",
    leg1="Mean F1",
    leg2="Mean FAR",
)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

COLORS = {
    "bg": RGBColor(248, 250, 252),
    "ink": RGBColor(15, 23, 42),
    "muted": RGBColor(71, 85, 105),
    "navy": RGBColor(30, 41, 59),
    "teal": RGBColor(15, 118, 110),
    "amber": RGBColor(180, 83, 9),
    "red": RGBColor(185, 28, 28),
    "blue": RGBColor(37, 99, 235),
    "line": RGBColor(203, 213, 225),
    "white": RGBColor(255, 255, 255),
}
FONT = "Microsoft JhengHei"
FONT_EN = "Times New Roman"


def set_bg(slide, color=COLORS["bg"]):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, text, x, y, w, h, size=24, bold=False, color=None, align="left", font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color or COLORS["ink"]
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.55, 0.28, 11.6, 0.55, size=25, bold=True)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.92), Inches(1.2), Inches(0.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS["teal"]
    bar.line.fill.background()
    if subtitle:
        add_text(slide, subtitle, 0.55, 1.02, 11.8, 0.35, size=11, color=COLORS["muted"])


def add_footer(slide, n):
    add_text(slide, f"TCSE 2026  |  CWRU-VAE  |  {n}", 10.4, 7.05, 2.4, 0.25, size=8, color=COLORS["muted"], align="right")


def add_bullets(slide, items, x, y, w, h, size=18):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.08)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        p.text = item
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = COLORS["ink"]
    return box


def add_card(slide, x, y, w, h, title, body, accent=COLORS["teal"]):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["white"]
    shape.line.color.rgb = COLORS["line"]
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    add_text(slide, title, x + 0.25, y + 0.15, w - 0.35, 0.35, size=15, bold=True, color=accent)
    add_text(slide, body, x + 0.25, y + 0.58, w - 0.35, h - 0.7, size=13)


def add_metric(slide, x, y, w, h, value, label, color=COLORS["teal"]):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["white"]
    shape.line.color.rgb = COLORS["line"]
    add_text(slide, value, x + 0.08, y + 0.16, w - 0.16, 0.45, size=24, bold=True, color=color, align="center")
    add_text(slide, label, x + 0.1, y + 0.72, w - 0.2, 0.45, size=11, color=COLORS["muted"], align="center")


def add_picture_fit(slide, path, x, y, w, h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_table(slide, rows, cols, x, y, w, h, data, font_size=11):
    table = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h)).table
    for c in range(cols):
        table.columns[c].width = Inches(w / cols)
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = str(data[r][c])
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    run.font.name = FONT
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = COLORS["ink"]
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLORS["navy"]
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = COLORS["white"]
    return table


slides_notes = []


def new_slide(title, subtitle=None):
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, title, subtitle)
    add_footer(slide, len(prs.slides))
    return slide


s = prs.slides.add_slide(blank)
set_bg(s, RGBColor(241, 245, 249))
add_text(s, "面向工業設備健康監測之\n可重現無監督異常偵測軟體流程", 0.7, 0.75, 8.7, 1.3, size=30, bold=True)
add_text(s, "以 VAE 軸承振動分析為例", 0.72, 2.18, 7.5, 0.45, size=20, color=COLORS["teal"], bold=True)
add_text(s, "A Reproducible Unsupervised Anomaly Detection Pipeline for Industrial Health Monitoring", 0.72, 2.75, 8.7, 0.42, size=14, color=COLORS["muted"], font=FONT_EN)
add_picture_fit(s, ROOT / "results/plots/01_time_domain.png", 8.35, 0.8, 4.2, 2.8)
add_card(s, 0.72, 4.55, 3.25, 1.05, "報告者", "張家軒\n國立勤益科技大學 人工智慧應用工程系", COLORS["teal"])
add_card(s, 4.25, 4.55, 3.25, 1.05, "指導教授", "張俊隆 教授\nNational Chin-Yi University of Technology", COLORS["blue"])
add_card(s, 7.78, 4.55, 3.25, 1.05, "研究定位", "可重現流程、正常資料訓練、跨工況閾值校準", COLORS["amber"])
add_footer(s, 1)
slides_notes.append("開場：本研究重點不是宣稱 VAE 在 CWRU 上突破效能，而是建立可重現且可檢驗部署限制的異常偵測流程。")

s = new_slide("報告大綱")
add_bullets(s, [
    "研究背景：為什麼工業異常偵測不能只看 AUC",
    "方法流程：CWRU 資料、VAE、正常資料訓練",
    "評審補強：load-wise split、threshold calibration、noise robustness",
    "主要發現：模型排序能力穩定，但固定閾值跨負載不穩",
    "結論：可重現流程 + 實務部署校準策略",
], 1.0, 1.35, 11.2, 4.6, size=22)
slides_notes.append("簡短說明報告路線：先講問題，再講方法，最後把評審要求的實驗整合成核心發現。")

s = new_slide("研究動機：真實工廠資料不是平衡分類題")
add_card(s, 0.75, 1.35, 3.6, 1.65, "故障資料稀少", "設備大多數時間處於健康狀態，故障事件少且昂貴。", COLORS["red"])
add_card(s, 4.85, 1.35, 3.6, 1.65, "標注成本高", "需要專家判讀，且很難事先蒐集所有故障模式。", COLORS["amber"])
add_card(s, 8.95, 1.35, 3.6, 1.65, "部署條件會變", "負載、轉速、感測雜訊與機台狀態都可能改變。", COLORS["blue"])
add_text(s, "因此，本研究採用無監督異常偵測：只用正常資料學習健康狀態，推論時以偏離正常分布程度作為異常分數。", 1.1, 4.0, 11.2, 1.0, size=22, bold=True, align="center")
slides_notes.append("強調工業情境和一般監督式分類不同：真正重要的是正常資料訓練與部署穩定性。")

s = new_slide("研究貢獻與評審回應")
add_bullets(s, [
    "可重現開源流程：自動下載、切窗、訓練、評估與視覺化",
    "正常資料訓練：不使用故障標籤學習分類邊界",
    "更嚴格泛化驗證：新增 load-wise split",
    "實務部署分析：補充 FAR、MR、PR-AUC、F1 與閾值校準",
    "模型定位修正：AE/VAE 對照，避免過度宣稱 VAE 效能突破",
], 0.8, 1.25, 6.6, 5.1, size=17)
add_card(s, 7.75, 1.45, 4.6, 1.15, "原始疑慮", "CWRU 上多數方法 AUC 皆接近 1.0，可能高估模型效能。", COLORS["red"])
add_card(s, 7.75, 3.0, 4.6, 1.15, "補強策略", "從「模型滿分」轉為檢驗「跨負載與閾值部署穩定性」。", COLORS["teal"])
add_card(s, 7.75, 4.55, 4.6, 1.15, "論文定位", "可重現流程與實務分析，而非單一模型效能突破。", COLORS["blue"])
slides_notes.append("這張回答審稿重點：我們不是硬拗 VAE 比別人強，而是承認 CWRU 容易，並補更嚴格的實驗。")

s = new_slide("軟體流程：從資料到部署指標")
steps = [("1", "CWRU\n資料下載"), ("2", "滑動視窗\n1024 / 512"), ("3", "正常資料\n訓練 VAE"), ("4", "重建誤差\n異常分數"), ("5", "load-wise\n泛化驗證"), ("6", "閾值校準\nFAR / MR")]
for i, (num, txt) in enumerate(steps):
    x = 0.55 + i * 2.08
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.0), Inches(1.55), Inches(1.25))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["white"]
    shape.line.color.rgb = COLORS["line"]
    add_text(s, num, x + 0.12, 2.12, 0.35, 0.3, size=16, bold=True, color=COLORS["teal"], align="center")
    add_text(s, txt, x + 0.15, 2.48, 1.25, 0.55, size=13, bold=True, align="center")
    if i < len(steps) - 1:
        add_text(s, "→", x + 1.62, 2.42, 0.38, 0.35, size=22, bold=True, color=COLORS["muted"], align="center")
add_text(s, "設計原則：流程模組化，模型可替換，評估指標不只看 AUC，也看部署時最敏感的誤報與漏報。", 1.15, 4.25, 11.0, 0.8, size=21, bold=True, align="center")
slides_notes.append("把研究講成可重現 pipeline，這是軟體工程研討會最容易接受的貢獻角度。")

s = new_slide("資料集與前處理")
add_picture_fit(s, ROOT / "results/plots/04_class_distribution.png", 0.75, 1.35, 5.15, 3.8)
add_card(s, 6.3, 1.25, 5.8, 1.1, "資料集", "CWRU Bearing Dataset，12 kHz；正常軸承與內圈、滾動體、外圈故障。", COLORS["blue"])
add_card(s, 6.3, 2.65, 5.8, 1.1, "切窗設定", "1,024 點視窗、512 點步長，50% 重疊，共 11,832 個片段。", COLORS["teal"])
add_card(s, 6.3, 4.05, 5.8, 1.1, "不平衡模擬", "保留全部正常樣本，只保留 10% 故障樣本，形成 79.5% : 20.5%。", COLORS["amber"])
slides_notes.append("說明資料不是一般平衡分類；我們刻意保留不平衡情境，較接近工廠健康監測。")

s = new_slide("VAE 異常偵測概念")
add_text(s, "只用正常訊號訓練", 0.8, 1.45, 2.3, 0.4, size=17, bold=True, color=COLORS["teal"], align="center")
add_text(s, "Encoder\nμ, log σ²", 3.55, 1.45, 2.0, 0.55, size=17, bold=True, color=COLORS["blue"], align="center")
add_text(s, "Latent z\nKL regularization", 6.15, 1.45, 2.2, 0.55, size=17, bold=True, color=COLORS["amber"], align="center")
add_text(s, "Decoder\n重建訊號", 8.95, 1.45, 2.0, 0.55, size=17, bold=True, color=COLORS["blue"], align="center")
for x in [0.9, 3.65, 6.25, 9.05]:
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.2), Inches(1.8), Inches(1.0))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["white"]
    shape.line.color.rgb = COLORS["line"]
for x in [2.75, 5.55, 8.45]:
    add_text(s, "→", x, 2.45, 0.45, 0.35, size=26, bold=True, color=COLORS["muted"], align="center")
add_text(s, "異常分數 = 平均重建 MSE", 3.7, 4.15, 5.8, 0.55, size=25, bold=True, color=COLORS["red"], align="center")
add_bullets(s, ["正常訊號：模型能重建，誤差低", "故障訊號：偏離正常潛在流形，誤差高", "本研究重點：異常分數排序 + 跨負載閾值穩定性"], 1.35, 5.15, 10.8, 1.15, size=17)
slides_notes.append("VAE 不必講太多數學，報告時聚焦：只用正常資料學習，重建誤差就是異常分數。")

s = new_slide("視窗層級基準：CWRU 高度可分")
add_picture_fit(s, ROOT / "results/plots/08_reconstruction_error.png", 0.75, 1.25, 5.25, 3.75)
add_picture_fit(s, ROOT / "results/plots/09_roc_curve.png", 6.25, 1.25, 4.8, 3.75)
add_metric(s, 1.25, 5.45, 2.4, 0.95, "1.0000", "AUC-ROC")
add_metric(s, 4.1, 5.45, 2.4, 0.95, "5.2×", "故障/正常重建誤差")
add_metric(s, 6.95, 5.45, 2.4, 0.95, "1.0000", "AE / VAE 皆滿分", COLORS["amber"])
add_text(s, "重點：滿分 AUC 代表資料容易分離，不足以證明實務部署穩定。", 0.95, 6.65, 11.2, 0.35, size=16, bold=True, color=COLORS["red"], align="center")
slides_notes.append("這張很重要：承認原始結果漂亮，但也指出它不足，因此自然帶到 load-wise split。")

s = new_slide("Load-wise Split：排序穩定，但固定閾值失效")
add_picture_fit(s, ASSET_DIR / "loadwise_far.png", 0.75, 1.25, 6.0, 3.6)
load_table = [["Load", "AUC", "PR-AUC", "F1", "FAR"]] + [[str(int(r.heldout_load)), f"{r.roc_auc:.4f}", f"{r.pr_auc:.4f}", f"{r.f1 * 100:.1f}%", f"{r.false_alarm_rate * 100:.1f}%"] for r in load_df.itertuples()]
add_table(s, len(load_table), 5, 7.05, 1.35, 5.35, 2.55, load_table, font_size=10)
add_card(s, 7.15, 4.45, 5.15, 1.3, "關鍵發現", "所有 load 的 AUC / PR-AUC 皆為 1.0000，但 load 3 的 FAR 達 99.16%。", COLORS["red"])
add_text(s, "AUC 會說「排序正確」，FAR 才會說「部署會不會一直誤報」。", 0.9, 6.35, 11.5, 0.45, size=18, bold=True, align="center")
slides_notes.append("說明 load-wise 是保留一個負載當未見測試；結果排序沒問題，但 P95 閾值跨負載不穩。")

s = new_slide("Threshold Calibration：少量目標正常資料大幅降低誤報")
add_picture_fit(s, ASSET_DIR / "calibration.png", 0.75, 1.2, 6.6, 4.1)
add_metric(s, 7.8, 1.45, 2.2, 1.0, "28.55%", "Val P95 平均 FAR", COLORS["red"])
add_metric(s, 10.15, 1.45, 2.2, 1.0, "1.14%", "Target P99 平均 FAR", COLORS["teal"])
add_metric(s, 7.8, 3.0, 2.2, 1.0, "78.89%", "Val P95 平均 F1", COLORS["amber"])
add_metric(s, 10.15, 3.0, 2.2, 1.0, "98.05%", "Target P99 平均 F1", COLORS["teal"])
add_card(s, 7.8, 4.65, 4.55, 1.05, "部署建議", "上線前蒐集少量目標負載健康資料，用 P99 校準閾值。", COLORS["teal"])
slides_notes.append("把校準講成部署流程：模型不用重訓，只要少量 target normal calibration 就能大幅降低 FAR。")

s = new_slide("Noise Robustness：CWRU 乾淨，低 SNR 會暴露風險")
add_picture_fit(s, ASSET_DIR / "noise.png", 0.75, 1.25, 6.8, 4.0)
add_card(s, 8.05, 1.35, 4.1, 1.0, "30 / 20 dB", "表現接近 clean setting，F1 約 78%。", COLORS["teal"])
add_card(s, 8.05, 2.75, 4.1, 1.0, "10 dB", "F1 降至 61.38%，FAR 上升至 52.90%。", COLORS["amber"])
add_card(s, 8.05, 4.15, 4.1, 1.0, "5 dB", "固定閾值幾乎失效，FAR 達 100%。", COLORS["red"])
add_text(s, "結論：真實工廠部署需搭配濾波、噪音感知校準或自適應閾值。", 0.9, 6.35, 11.4, 0.4, size=18, bold=True, align="center")
slides_notes.append("強調這不是模型失敗，而是說明 CWRU 乾淨；低 SNR 下固定閾值不可靠。")

s = new_slide("AE vs VAE：誠實定位模型價值")
ae_mean = ae_df[["roc_auc", "pr_auc", "f1", "false_alarm_rate", "miss_rate"]].mean()
vae_mean = load_df[["roc_auc", "pr_auc", "f1", "false_alarm_rate", "miss_rate"]].mean()
model_table = [
    ["Model", "AUC", "PR-AUC", "Mean F1", "Mean FAR", "Mean MR"],
    ["VAE", f"{vae_mean.roc_auc:.4f}", f"{vae_mean.pr_auc:.4f}", f"{vae_mean.f1 * 100:.2f}%", f"{vae_mean.false_alarm_rate * 100:.2f}%", f"{vae_mean.miss_rate * 100:.2f}%"],
    ["AE", f"{ae_mean.roc_auc:.4f}", f"{ae_mean.pr_auc:.4f}", f"{ae_mean.f1 * 100:.2f}%", f"{ae_mean.false_alarm_rate * 100:.2f}%", f"{ae_mean.miss_rate * 100:.2f}%"],
]
add_table(s, 3, 6, 0.9, 1.45, 11.5, 1.45, model_table, font_size=12)
add_card(s, 1.0, 3.55, 3.4, 1.35, "排序指標", "AE 與 VAE 的 AUC / PR-AUC 皆達 1.0000。", COLORS["blue"])
add_card(s, 4.95, 3.55, 3.4, 1.35, "閾值指標", "兩者 FAR 接近，AE 略低；不應宣稱 VAE 顯著優越。", COLORS["amber"])
add_card(s, 8.9, 3.55, 3.4, 1.35, "VAE 價值", "作為具機率潛在空間的可替換模型元件。", COLORS["teal"])
add_text(s, "研究定位從「VAE 效能突破」修正為「可重現管線與部署限制分析」。", 1.0, 6.1, 11.3, 0.45, size=19, bold=True, align="center")
slides_notes.append("這張可回應評審：我們沒有硬說 VAE 比 AE 強，而是誠實呈現模型差異。")

s = new_slide("核心發現：真正的問題不是分離，而是部署閾值")
add_card(s, 0.85, 1.25, 3.7, 1.6, "1. CWRU 高度可分", "視窗層級與 load-wise 下，AUC / PR-AUC 仍可滿分。", COLORS["blue"])
add_card(s, 4.85, 1.25, 3.7, 1.6, "2. 固定閾值不穩", "load 3 的 FAR 顯示跨負載分布轉移會造成高誤報。", COLORS["red"])
add_card(s, 8.85, 1.25, 3.7, 1.6, "3. 校準很有效", "少量目標正常資料能將 FAR 從 28.55% 降至 1.14%。", COLORS["teal"])
add_card(s, 2.65, 3.65, 3.7, 1.6, "4. 噪音仍是挑戰", "10 / 5 dB 下固定閾值明顯失效，需要自適應策略。", COLORS["amber"])
add_card(s, 6.95, 3.65, 3.7, 1.6, "5. 模型可替換", "AE/VAE 差異有限，流程本身才是可延伸貢獻。", COLORS["navy"])
slides_notes.append("總結三個層次：資料可分、閾值漂移、部署校準。這是報告最核心的 takeaway。")

s = new_slide("結論與未來工作")
add_bullets(s, [
    "本文提出一套可重現的無監督工業異常偵測流程。",
    "VAE 在 CWRU 上能穩定分離正常與故障訊號，但滿分 AUC 不等於部署可靠。",
    "Load-wise split 顯示跨負載閾值漂移是實務風險。",
    "Threshold calibration 可有效降低 false alarm rate。",
    "未來將延伸至 IMS、FEMTO-ST / PRONOSTIA、Paderborn 等更具挑戰資料集。",
], 0.9, 1.3, 8.2, 4.9, size=18)
add_card(s, 9.45, 1.55, 2.75, 1.0, "GitHub", "github.com/charmi582/CWRU-VAE", COLORS["teal"])
add_card(s, 9.45, 3.0, 2.75, 1.0, "投稿定位", "可重現流程\n正常資料訓練\n部署校準分析", COLORS["blue"])
add_text(s, "謝謝聆聽", 9.45, 5.0, 2.75, 0.55, size=25, bold=True, align="center")
slides_notes.append("結尾再次呼應：流程可重現、評估更嚴格、未來做跨資料集與部署校準。")

pptx_path = OUT_DIR / "TCSE2026_CWRU_VAE_presentation.pptx"
prs.save(pptx_path)

notes_path = OUT_DIR / "TCSE2026_CWRU_VAE_speaker_notes.md"
with notes_path.open("w", encoding="utf-8") as f:
    f.write("# TCSE2026 CWRU-VAE 研討會報告講稿大綱\n\n")
    for i, note in enumerate(slides_notes, 1):
        f.write(f"## Slide {i}\n{note}\n\n")

print(pptx_path)
print(notes_path)
print(f"slides {len(prs.slides)}")
